import os
from typing import Dict, Any, List
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from backend.app.core.config import settings

class ExportService:
    @staticmethod
    def generate_docx(meeting_data: Dict[str, Any], output_path: str) -> str:
        doc = Document()
        
        # Styles
        title_p = doc.add_paragraph()
        title_run = title_p.add_run(meeting_data.get("title", "Minutes of Meeting"))
        title_run.font.size = Pt(22)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(26, 54, 93) # Navy
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph(f"Date: {meeting_data.get('date', 'N/A')} | Duration: {meeting_data.get('duration', 0)}s | Type: {meeting_data.get('meeting_type', 'LIVE')}")
        participants = ", ".join([p.get("name", "") for p in meeting_data.get("participants", [])])
        doc.add_paragraph(f"Participants: {participants}")
        doc.add_heading("1. Executive Summary", level=1)
        doc.add_paragraph(meeting_data.get("summary", "No summary provided."))

        # Topics
        topics = meeting_data.get("topics", [])
        if topics:
            doc.add_heading("2. Topics Discussed", level=1)
            for t in topics:
                p = doc.add_paragraph(style='List Bullet')
                p.add_run(f"{t.get('topic_name')}: ").bold = True
                p.add_run(t.get('summary', ''))

        # Decisions
        decisions = meeting_data.get("decisions", [])
        if decisions:
            doc.add_heading("3. Key Decisions", level=1)
            for d in decisions:
                p = doc.add_paragraph(style='List Bullet')
                p.add_run(d.get('decision', '')).bold = True
                if d.get('evidence'):
                    ev = doc.add_paragraph(style='List Continue')
                    evidence_text = d.get('evidence', '')
                    ev_run = ev.add_run(f"Evidence: \"{evidence_text}\"")
                    ev_run.font.italic = True

        # Action Items Table
        actions = meeting_data.get("action_items", [])
        if actions:
            doc.add_heading("4. Action Items", level=1)
            table = doc.add_table(rows=1, cols=5)
            table.style = 'Light Shading Accent 1'
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = 'Task'
            hdr_cells[1].text = 'Owner'
            hdr_cells[2].text = 'Deadline'
            hdr_cells[3].text = 'Status'
            hdr_cells[4].text = 'Evidence'
            
            for a in actions:
                row_cells = table.add_row().cells
                row_cells[0].text = a.get("task", "")
                row_cells[1].text = a.get("owner", "NEEDS_REVIEW")
                row_cells[2].text = a.get("deadline", "NEEDS_REVIEW")
                row_cells[3].text = a.get("status", "PENDING")
                row_cells[4].text = a.get("evidence", "")

        # Unresolved Issues
        unresolved = meeting_data.get("unresolved_issues", [])
        if unresolved:
            doc.add_heading("5. Unresolved Issues", level=1)
            for u in unresolved:
                p = doc.add_paragraph(style='List Bullet')
                p.add_run(u.get('issue', ''))

        doc.save(output_path)
        return output_path

    @staticmethod
    def generate_pdf(meeting_data: Dict[str, Any], output_path: str) -> str:
        pdf = SimpleDocTemplate(output_path, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
        styles = getSampleStyleSheet()
        story = []

        title_style = ParagraphStyle('TitleStyle', parent=styles['Heading1'], fontSize=20, textColor=colors.HexColor('#1E3A8A'), spaceAfter=12)
        h1_style = ParagraphStyle('H1Style', parent=styles['Heading2'], fontSize=14, textColor=colors.HexColor('#0F172A'), spaceBefore=10, spaceAfter=6)
        body_style = styles['Normal']

        story.append(Paragraph(meeting_data.get("title", "Minutes of Meeting"), title_style))
        meta_text = f"<b>Date:</b> {meeting_data.get('date', 'N/A')} | <b>Duration:</b> {meeting_data.get('duration', 0)}s"
        story.append(Paragraph(meta_text, body_style))
        story.append(Spacer(1, 10))

        story.append(Paragraph("1. Executive Summary", h1_style))
        story.append(Paragraph(meeting_data.get("summary", "No summary."), body_style))
        story.append(Spacer(1, 10))

        decisions = meeting_data.get("decisions", [])
        if decisions:
            story.append(Paragraph("2. Key Decisions", h1_style))
            for d in decisions:
                story.append(Paragraph(f"• <b>{d.get('decision')}</b>", body_style))
                if d.get("evidence"):
                    ev_t = d.get('evidence', '')
                    story.append(Paragraph(f"<i>Evidence: \"{ev_t}\"</i>", body_style))
            story.append(Spacer(1, 10))

        actions = meeting_data.get("action_items", [])
        if actions:
            story.append(Paragraph("3. Action Items", h1_style))
            table_data = [["Task", "Owner", "Deadline", "Status"]]
            for a in actions:
                table_data.append([
                    Paragraph(a.get("task", ""), body_style),
                    Paragraph(a.get("owner", "NEEDS_REVIEW"), body_style),
                    Paragraph(a.get("deadline", "NEEDS_REVIEW"), body_style),
                    Paragraph(a.get("status", "PENDING"), body_style)
                ])
            t = Table(table_data, colWidths=[240, 100, 100, 80])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#F1F5F9')),
                ('TEXTCOLOR', (0,0), (-1,0), colors.HexColor('#0F172A')),
                ('ALIGN', (0,0), (-1,-1), 'LEFT'),
                ('BOTTOMPADDING', (0,0), (-1,-1), 6),
                ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
            ]))
            story.append(t)
            story.append(Spacer(1, 10))

        pdf.build(story)
        return output_path

export_service = ExportService()
